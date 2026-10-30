"""
SmartMatch AI - Document Text Extractor
========================================
This module handles extracting raw text from various document formats:
- PDF (.pdf) using PyMuPDF (fitz) with fallback to pypdf
- Word (.docx) using python-docx
- Plain Text (.txt) using UTF-8/Latin-1 decoding

Beginner-Friendly Explanation:
Before we can calculate text similarity, we need to extract the readable
text from whatever file format the user uploaded, and break it down
into paragraphs so we can compare each paragraph with the user's question.
"""

import io
import re
from typing import List, Dict, Any, Tuple

# Try importing PyMuPDF (fitz), with fallback to pypdf
try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts all textual content from a PDF file using PyMuPDF (fitz),
    falling back to pypdf if PyMuPDF is not available.
    """
    extracted_text = []

    # Method 1: PyMuPDF (fastest & most accurate layout extraction)
    if fitz is not None:
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text")
                if text and text.strip():
                    extracted_text.append(text.strip())
            doc.close()
            if extracted_text:
                return "\n\n".join(extracted_text)
        except Exception as e:
            # Fall back to method 2 if PyMuPDF encounters any parsing glitch
            pass

    # Method 2: pypdf fallback
    if pypdf is not None:
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                text = page.extract_text()
                if text and text.strip():
                    extracted_text.append(text.strip())
            if extracted_text:
                return "\n\n".join(extracted_text)
        except Exception as e:
            raise ValueError(f"Could not read PDF file: {str(e)}")

    if not extracted_text:
        raise ValueError("Could not extract any readable text from this PDF file.")

    return "\n\n".join(extracted_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from Microsoft Word (.docx) documents paragraph by paragraph.
    """
    if docx is None:
        raise ImportError("python-docx is required to parse .docx files.")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        
        # Also check tables inside the Word document for tabular data
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        if not paragraphs:
            raise ValueError("The Word document is empty or contains no readable text.")

        return "\n\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Could not read Word document (.docx): {str(e)}")


def extract_text_from_txt(file_bytes: bytes) -> str:
    """
    Extracts text from a plain text (.txt) file with robust encoding detection.
    """
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            text = file_bytes.decode(encoding)
            if text.strip():
                return text.strip()
        except UnicodeDecodeError:
            continue

    raise ValueError("Could not decode the text file with supported encodings (UTF-8, Latin-1).")


def extract_document_text(file_bytes: bytes, filename: str) -> str:
    """
    Main extraction dispatcher based on file extension.
    Supported extensions: .pdf, .docx, .txt
    """
    filename_lower = filename.lower()
    
    if filename_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif filename_lower.endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file type for '{filename}'. Please upload a .pdf, .docx, or .txt file."
        )


def split_into_paragraphs(raw_text: str, min_words: int = 5) -> List[Dict[str, Any]]:
    """
    Splits the extracted document text into logical paragraphs or semantic chunks.
    Filters out noise, short headers, page numbers, and trivial lines.
    
    Returns a list of paragraph objects with:
    - index: 1-based paragraph sequence number
    - text: cleaned paragraph text
    - word_count: number of words in this paragraph
    """
    # Normalize line breaks
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    
    # Split primarily by double newlines (standard paragraphs)
    raw_blocks = re.split(r"\n\s*\n+", text)
    
    paragraphs = []
    current_index = 1
    
    for block in raw_blocks:
        # Collapse multiple internal spaces and single line breaks into a clean flow
        cleaned_block = re.sub(r"\s+", " ", block).strip()
        
        words = cleaned_block.split()
        # Keep paragraphs that have at least `min_words` (ignores page numbers like "1", "- 2 -")
        if len(words) >= min_words:
            paragraphs.append({
                "index": current_index,
                "text": cleaned_block,
                "word_count": len(words)
            })
            current_index += 1

    # Fallback: if document had no double line breaks, split by period/sentences
    if len(paragraphs) <= 1 and len(raw_text.split()) > 40:
        sentences = re.split(r"(?<=[.!?])\s+", text)
        chunk = []
        chunk_words = 0
        current_index = 1
        paragraphs = []
        
        for s in sentences:
            s_clean = s.strip()
            if not s_clean:
                continue
            words = s_clean.split()
            chunk.append(s_clean)
            chunk_words += len(words)
            
            # Form paragraph chunks of ~35-70 words
            if chunk_words >= 35:
                combined = " ".join(chunk)
                paragraphs.append({
                    "index": current_index,
                    "text": combined,
                    "word_count": chunk_words
                })
                current_index += 1
                chunk = []
                chunk_words = 0
                
        if chunk:
            paragraphs.append({
                "index": current_index,
                "text": " ".join(chunk),
                "word_count": chunk_words
            })

    return paragraphs
