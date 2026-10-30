import os
import io
from pathlib import Path
import docx
import pymupdf as fitz

BASE_DIR = Path(__file__).resolve().parent
SAMPLE_DIR = BASE_DIR / "sample_data"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

# 1. Diabetes Text File (Matches requested placeholder example)
diabetes_content = """# Clinical Overview of Diabetes Mellitus and Metabolic Pathophysiology

Diabetes mellitus is a chronic metabolic disorder characterized by elevated levels of blood glucose over a prolonged period. The disease results from either insufficient insulin production by pancreatic beta cells, or the body cells failing to respond appropriately to the insulin produced.

What are the main causes of diabetes?
The primary causes of diabetes depend on the specific classification. Type 1 diabetes is primarily caused by an autoimmune reaction where the immune system mistakenly attacks and destroys insulin-producing beta cells in the islet of Langerhans within the pancreas. Genetic susceptibility, human leukocyte antigen (HLA) genotypes, and environmental triggers such as viral infections are significant contributing factors.

In contrast, Type 2 diabetes develops from a combination of chronic insulin resistance and progressive secretory deficiency. The predominant risk factors and root causes include prolonged obesity, physical inactivity, atherogenic diets high in refined sugars and saturated fats, advancing age, and strong familial genetic predispositions. Chronic low-grade systemic inflammation and visceral adiposity disrupt insulin signaling pathways in hepatic, muscular, and adipose tissues.

Gestational diabetes occurs during pregnancy when placental hormones impair maternal insulin sensitivity, causing blood glucose levels to elevate beyond normal thresholds.

Management and therapeutic interventions include regular physical aerobic activity, balanced carbohydrate-controlled nutrition, continuous glucose monitoring, oral antihyperglycemic pharmacotherapy such as metformin and SGLT2 inhibitors, and exogenous subcutaneous insulin replacement when indicated.
"""

with open(SAMPLE_DIR / "diabetes_causes.txt", "w", encoding="utf-8") as f:
    f.write(diabetes_content.strip())
print("Created diabetes_causes.txt")

# 2. AI Overview Text File
ai_content = """# Fundamentals of Artificial Intelligence and Machine Learning

Artificial intelligence (AI) is the simulation of human intelligence processes by computer systems. These processes include the acquisition of information, the application of reasoning rules to reach approximate or definite conclusions, and self-correction.

Natural Language Processing (NLP) is a specialized branch of artificial intelligence focused on enabling computers to understand, interpret, and manipulate human language. Core NLP tasks include text classification, sentiment analysis, named entity recognition, and semantic similarity measurement using vector space models.

Term Frequency-Inverse Document Frequency (TF-IDF) is a foundational statistical method in NLP. It evaluates how important a word is to a document within a larger collection or corpus. Words that appear frequently in a specific paragraph but rarely across the rest of the document receive the highest mathematical weights.

Cosine similarity evaluates the orientation rather than magnitude between two multi-dimensional numerical vectors. By calculating the cosine of the angle between query and document vectors, systems can retrieve the most relevant passages regardless of paragraph length.
"""

with open(SAMPLE_DIR / "artificial_intelligence.txt", "w", encoding="utf-8") as f:
    f.write(ai_content.strip())
print("Created artificial_intelligence.txt")

# 3. Climate Change Word (.docx) Document
doc = docx.Document()
doc.add_heading("Global Climate Dynamics and Environmental Impact Assessment", level=1)
doc.add_paragraph("Climate change refers to long-term shifts in temperatures and weather patterns across the Earth's biosphere. While natural processes such as variations in the solar cycle influence the climate, human industrial activities have been the undisputed primary driver of global warming since the mid-19th century.")
doc.add_heading("Primary Drivers of Global Warming", level=2)
doc.add_paragraph("The combustion of fossil fuels—coal, oil, and gas—releases billions of metric tons of carbon dioxide and nitrous oxide into the atmosphere. These heat-trapping greenhouse gases blanket the Earth, preventing thermal radiation from escaping into space and elevating average global temperatures.")
doc.add_heading("Ecosystem and Biosphere Consequences", level=2)
doc.add_paragraph("Consequences of rapid climate change include intensifying droughts, water scarcity, catastrophic wildfires, rising global sea levels from melting polar ice sheets, catastrophic flooding, and the loss of biodiversity across marine coral reefs and terrestrial habitats.")
doc.add_heading("Mitigation and Renewable Transition", level=2)
doc.add_paragraph("Transitioning from fossil fuels to clean renewable energy sources such as utility-scale solar photovoltaics, wind turbines, geothermal energy, and green hydrogen represents the cornerstone of international decarbonization commitments under the Paris Agreement.")
doc.save(str(SAMPLE_DIR / "climate_change.docx"))
print("Created climate_change.docx")

# 4. Quantum Computing PDF (.pdf) Document
pdf_doc = fitz.open()
page = pdf_doc.new_page(width=595, height=842) # A4 size

p1 = """Quantum Computing Foundations and Qubit Mechanics

Quantum computing is an emerging field of computer science that harnesses the collective properties of quantum states, such as superposition, interference, and entanglement, to perform calculations exponentially faster than classical computers for specific problem domains.

Superposition and Quantum Information Units
Unlike classical computing where information is encoded in binary bits that must be strictly 0 or 1, a quantum bit (qubit) can exist in a linear combination of both states simultaneously. This multi-state superposition enables quantum processors to evaluate vast computational search spaces concurrently.

Quantum Algorithms and Cryptographic Applications
Prominent quantum algorithms include Shor's algorithm for polynomial-time integer factorization—posing significant security implications for classical RSA cryptography—and Grover's algorithm for quadratic speedup in unstructured database search queries.

Fault Tolerance and Physical Qubit Decoherence
A major engineering hurdle in quantum computing is quantum decoherence and environmental thermal noise. To achieve fault-tolerant quantum computation, researchers are developing topological quantum codes and surface error correction lattices requiring thousands of physical qubits to encode a single stable logical qubit."""

rect = fitz.Rect(50, 50, 545, 792)
page.insert_textbox(rect, p1, fontsize=11, fontname="helv")
pdf_doc.save(str(SAMPLE_DIR / "quantum_computing.pdf"))
pdf_doc.close()
print("Created quantum_computing.pdf")
print("All sample files generated successfully!")
